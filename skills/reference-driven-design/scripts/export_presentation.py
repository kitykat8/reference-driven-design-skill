#!/usr/bin/env python3
"""
Presentation Exporter for Reference-Driven Design (`reference-driven-design` & `slides` skills)
Exports multi-slide HTML/SVG/CSS design systems into studio-grade:
  1. Vector 16:9 PDF (.pdf) with embedded Google Fonts, crisp SVGs, and zero margins.
  2. Widescreen 16:9 PowerPoint (.pptx) with Retina 2560x1440 slide visuals and extracted slide titles/notes.
  3. Contact Preview Sheet (_preview.jpg) so the agent can visually verify every page via `view_file`.
"""

import argparse
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


BROWSER_CANDIDATES = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]


def find_browser() -> str:
    for path in BROWSER_CANDIDATES:
        if os.path.exists(path):
            return path
    raise FileNotFoundError("Could not find Chrome or Edge browser for headless PDF rendering.")


def get_print_override_css(base_w: int = 800) -> str:
    base_h = int(round(base_w * 9 / 16))
    return f"""
<style id="pdf-print-override">
  @page {{
    size: {base_w}px {base_h}px;
    margin: 0 !important;
  }}
  @media print {{
    * {{
      -webkit-print-color-adjust: exact !important;
      print-color-adjust: exact !important;
      transition: none !important;
      animation: none !important;
    }}
    html, body {{
      width: {base_w}px !important;
      height: auto !important;
      margin: 0 !important;
      padding: 0 !important;
      overflow: visible !important;
      background: #0A0A0A !important;
    }}
    .deck-controls, .hud, nav, .controls, #controls {{
      display: none !important;
    }}
    .deck-container, .deck, main {{
      display: block !important;
      padding: 0 !important;
      margin: 0 !important;
      max-width: none !important;
      width: {base_w}px !important;
      height: auto !important;
      overflow: visible !important;
    }}
    .slide {{
      display: flex !important;
      position: relative !important;
      width: {base_w}px !important;
      height: {base_h}px !important;
      min-height: {base_h}px !important;
      max-height: {base_h}px !important;
      aspect-ratio: 16 / 9 !important;
      margin: 0 !important;
      page-break-after: always !important;
      break-after: page !important;
      page-break-inside: avoid !important;
      break-inside: avoid !important;
      overflow: hidden !important;
      box-shadow: none !important;
      transform: none !important;
    }}
    .slide:last-of-type {{
      page-break-after: auto !important;
      break-after: auto !important;
    }}
  }}
</style>
"""


def render_html_to_pdf(html_path: str, pdf_path: str, expected_slides: int = 0, base_w: int = 800) -> int:
    browser = find_browser()
    with open(html_path, "r", encoding="utf-8") as f:
        html_content = f.read()

    override_css = get_print_override_css(base_w)
    if "</head>" in html_content:
        patched_html = html_content.replace("</head>", override_css + "\n</head>", 1)
    else:
        patched_html = override_css + "\n" + html_content

    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8", dir=os.path.dirname(os.path.abspath(html_path))) as tmp:
        tmp.write(patched_html)
        tmp_html_path = tmp.name

    raw_pdf_path = pdf_path + ".raw.pdf"
    try:
        file_url = Path(tmp_html_path).resolve().as_uri()
        cmd = [
            browser,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            "--virtual-time-budget=4000",
            f"--print-to-pdf={os.path.abspath(raw_pdf_path)}",
            file_url,
        ]
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

        # Clean up any trailing blank pages using PyMuPDF
        doc = pymupdf.open(raw_pdf_path)
        if expected_slides > 0 and len(doc) > expected_slides:
            while len(doc) > expected_slides:
                doc.delete_page(len(doc) - 1)
        doc.save(pdf_path)
        page_count = len(doc)
        doc.close()
        return page_count
    finally:
        if os.path.exists(tmp_html_path):
            os.remove(tmp_html_path)
        if os.path.exists(raw_pdf_path):
            os.remove(raw_pdf_path)


def extract_slide_titles_from_html(html_path: str) -> list[str]:
    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()
    slides = re.split(r'<section[^>]*class="[^"]*slide[^"]*"[^>]*>', html)[1:]
    titles = []
    for idx, s in enumerate(slides, 1):
        m = re.search(r'<div class="h-display"[^>]*>(.*?)</div>', s, re.DOTALL)
        if m:
            clean = re.sub(r"<[^>]+>", " ", m.group(1))
            clean = re.sub(r"\s+", " ", clean).strip()
            titles.append(f"Slide {idx:02d}: {clean}")
        else:
            titles.append(f"Slide {idx:02d}")
    return titles


def export_pdf_to_pptx(pdf_path: str, pptx_path: str, slide_titles: list[str] | None = None, target_w: int = 2560) -> int:
    doc = pymupdf.open(pdf_path)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # blank layout

    with tempfile.TemporaryDirectory() as tmpdir:
        for i, page in enumerate(doc):
            scale = target_w / page.rect.width
            mat = pymupdf.Matrix(scale, scale)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            img_path = os.path.join(tmpdir, f"slide_{i+1:02d}.png")
            pix.save(img_path)

            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                img_path,
                Inches(0),
                Inches(0),
                width=prs.slide_width,
                height=prs.slide_height,
            )

            if slide_titles and i < len(slide_titles):
                notes_slide = slide.notes_slide
                notes_slide.notes_text_frame.text = slide_titles[i]

    prs.save(pptx_path)
    count = len(doc)
    doc.close()
    return count


def build_contact_sheet(pdf_path: str, preview_jpg_path: str, cols: int = 2) -> None:
    doc = pymupdf.open(pdf_path)
    thumbs = []
    thumb_w, thumb_h = 800, 450
    for page in doc:
        pix = page.get_pixmap(matrix=pymupdf.Matrix(thumb_w / page.rect.width, thumb_h / page.rect.height), alpha=False)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        thumbs.append(img)
    doc.close()

    rows = (len(thumbs) + cols - 1) // cols
    pad = 24
    header_h = 64
    sheet_w = cols * thumb_w + (cols + 1) * pad
    sheet_h = header_h + rows * thumb_h + (rows + 1) * pad

    sheet = Image.new("RGB", (sheet_w, sheet_h), (14, 14, 14))
    draw = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype(r"C:\Windows\Fonts\consolab.ttf", 20)
    except Exception:
        font = ImageFont.load_default()

    draw.text((pad, 20), f"EXPORTED PRESENTATION VERIFICATION SHEET // {len(thumbs)} SLIDES (16:9 PDF & PPTX)", fill=(239, 232, 82), font=font)

    for idx, im in enumerate(thumbs):
        r = idx // cols
        c = idx % cols
        x = pad + c * (thumb_w + pad)
        y = header_h + pad + r * (thumb_h + pad)
        sheet.paste(im, (x, y))
        draw.rectangle([x - 1, y - 1, x + thumb_w, y + thumb_h], outline=(239, 232, 82), width=1)

    sheet.save(preview_jpg_path, format="JPEG", quality=90)


def main():
    parser = argparse.ArgumentParser(description="Export HTML design-system slides to PDF and/or PPTX")
    parser.add_argument("--input", "-i", required=True, help="Path to input HTML presentation")
    parser.add_argument("--pdf", help="Output PDF file path (.pdf)")
    parser.add_argument("--pptx", help="Output PowerPoint file path (.pptx)")
    parser.add_argument("--preview", help="Output contact sheet preview JPG (.jpg)")
    parser.add_argument("--slides", type=int, default=0, help="Expected number of slides (trims extra blank pages)")
    parser.add_argument("--base-width", type=int, default=800, help="16:9 CSS design width for vector PDF layout (default: 800 -> 800x450)")
    args = parser.parse_args()

    if not args.pdf and not args.pptx:
        print("Error: Specify at least one of --pdf or --pptx", file=sys.stderr)
        sys.exit(1)

    # Always generate PDF first (either at args.pdf or a temp path if only PPTX was requested)
    cleanup_pdf = False
    pdf_target = args.pdf
    if not pdf_target:
        cleanup_pdf = True
        pdf_target = args.pptx + ".tmp.pdf"

    os.makedirs(os.path.dirname(os.path.abspath(pdf_target)), exist_ok=True)
    num_pages = render_html_to_pdf(args.input, pdf_target, expected_slides=args.slides, base_w=args.base_width)
    print(f"[OK] Rendered {num_pages}-page 16:9 vector PDF -> {pdf_target}")

    if args.pptx:
        os.makedirs(os.path.dirname(os.path.abspath(args.pptx)), exist_ok=True)
        titles = extract_slide_titles_from_html(args.input)
        num_pptx = export_pdf_to_pptx(pdf_target, args.pptx, slide_titles=titles, target_w=2560)
        print(f"[OK] Exported {num_pptx}-slide 16:9 widescreen PPTX -> {args.pptx}")

    if args.preview:
        os.makedirs(os.path.dirname(os.path.abspath(args.preview)), exist_ok=True)
        build_contact_sheet(pdf_target, args.preview)
        print(f"[OK] Generated verification preview sheet -> {args.preview}")

    if cleanup_pdf and os.path.exists(pdf_target):
        os.remove(pdf_target)


if __name__ == "__main__":
    main()
